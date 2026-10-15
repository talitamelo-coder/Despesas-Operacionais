"""
Converte a planilha "Fornecedores e Licença" em um arquivo de importação de licenças
(formato akross-licencas-importacao), restaurável em Licenças → Importar licenças.

- Aba Software = catálogo (produtos, custo, IOF/impostos, ficha da ferramenta).
- Abas por ferramenta = quem usa (atribuições). Microsoft traz os próprios planos.
- Planilha1 (faturas de telefonia) = só as linhas do mês mais recente.
- Cancelados = histórico.
- Pessoas unificadas pelo e-mail; sem e-mail, pelo nome (marcado para confirmar).
- Vínculo com contrato sugerido pelo fornecedor, sempre marcado para confirmar.
Não altera contratos.

Uso: python3 -I converter_licencas.py <licencas.xlsx> <backup-contratos.json> <saida.json> <relatorio.xlsx>
"""
import sys, json, re, unicodedata, datetime, warnings, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

warnings.filterwarnings("ignore")
xlsx, backup_path, out_path, rel_path = sys.argv[1:5]
wb = openpyxl.load_workbook(xlsx, data_only=True)
backup = json.load(open(backup_path, encoding="utf-8"))["estado"]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


def texto(v):
    if v is None:
        return ""
    s = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()
    return "" if s.lower() in ("none", "-", "--", "n/a") else s


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


PART = {"de", "da", "do", "das", "dos", "e"}


def titulo(s):
    return " ".join(w if w in PART else w.capitalize() for w in s.lower().split())


# ------------------------------------------------------------------ vínculo com contratos (backup)
GENERICOS = {"ltda", "s", "a", "sa", "me", "inc", "servicos", "tecnologia", "brasil", "do", "de", "da", "the"}


def chave_fornecedor(nome):
    return next((w for w in norm(nome).split() if len(w) >= 3 and w not in GENERICOS), norm(nome))


forn_por_id = {f["fornecedor_id"]: f for f in backup["fornecedores"]}


def contrato_sugerido(fornecedor_nome):
    k = chave_fornecedor(fornecedor_nome)
    ids = {fid for fid, f in forn_por_id.items() if k in norm(f"{f['razaoSocial']} {f.get('nomeFantasia', '')}").split()}
    cands = [c for c in backup["contratos"] if c["fornecedorId"] in ids and c["status"] == "Vigente"]
    return (cands[0]["contrato_id"], cands[0]["codigo"]) if len(cands) == 1 else (None, None)


# ------------------------------------------------------------------ colaboradores
colabs = {}  # chave -> dict
por_nome = {}  # nome normalizado -> chave

COMPARTILHADA = re.compile(r"\b(admin|administrador|grupo|sdp|apps|wifi|data ads|plantao|mesa|financeiro|negocios|office\d*|suporte|contato|desativado)\b")


def local(email):
    e = texto(email).lower()
    return e.split("@")[0] if e else ""


def registrar_colaborador(nome, email="", area="", centro=""):
    nome, email = texto(nome), texto(email).lower()
    if not nome and not email:
        return None
    lp = local(email)
    chave = f"e:{lp}" if lp and "." in lp or (lp and "@" in email) else None
    if not chave:
        n = norm(nome)
        chave = por_nome.get(n)
        if not chave:
            # Primeiro + último nome iguais a alguém já conhecido.
            partes = n.split()
            if len(partes) >= 2:
                cand = {k for nn, k in por_nome.items() if nn.split()[:1] == partes[:1] and nn.split()[-1:] == partes[-1:]}
                if len(cand) == 1:
                    chave = cand.pop()
        if not chave:
            chave = f"n:{n}"
    c = colabs.setdefault(chave, {"chave": chave, "nomes": [], "email": "", "area": "", "centroCusto": "", "semEmail": True})
    if nome and nome.lower() != lp:
        c["nomes"].append(nome)
    elif nome:
        c["nomes"].append(titulo(nome.replace(".", " ")))
    if "@" in email and (not c["email"] or (c["email"].endswith("mobicare.com.br") and email.endswith("akross.com.br"))):
        c["email"] = email
    if "@" in email:
        c["semEmail"] = False
    if area and not c["area"]:
        c["area"] = area
    if centro and not c["centroCusto"]:
        c["centroCusto"] = centro
    for nm in c["nomes"]:
        por_nome.setdefault(norm(nm), chave)
    return chave


# ------------------------------------------------------------------ produtos
produtos = {}  # chave -> produto
atribuicoes = []
relatorio = []

ws = wb["Software"]
rows = list(ws.iter_rows(values_only=True))
# PTAX implícita nos totais da planilha (Valor Total Ano BRL ÷ valor em USD no ano).
ptax = None
for r in rows[5:]:
    if r[1] and texto(r[5]) == "USD" and num(r[6]) and num(r[4]) and num(r[8]):
        anual_usd = num(r[6]) * num(r[4]) * (12 if texto(r[3]) == "Mensal" else 1)
        ptax = round(num(r[8]) / anual_usd, 4)
        break
contagem_forn = collections.Counter(texto(r[1]) for r in rows[5:] if r[1])


def novo_produto(chave, **kw):
    p = {"chave": chave, "situacao": "Ativo", "categoria": "Software", "moeda": "BRL", "periodicidade": "Mensal", "quantidadeContratada": 0, "valorUnitario": 0, "confirmacoesPendentes": []}
    p.update({k: v for k, v in kw.items() if v not in (None, "")})
    produtos[chave] = p
    return p


for r in rows[5:]:
    forn = texto(r[1])
    if not forn:
        continue
    util = texto(r[2])
    multiplo = contagem_forn[forn] > 1
    nome = util if multiplo and util else forn
    chave = f"sw:{norm(forn)}:{norm(nome)[:40]}"
    if chave in produtos:
        chave += ":2"
    moeda = "USD" if texto(r[5]) == "USD" else "BRL"
    venc = r[13]
    o, pimp = num(r[14]), num(r[15])
    obs = []
    if num(r[16]):
        obs.append(f"Comparativo Nexforce: R$ {num(r[16]):,.2f}/ano")
    if num(r[17]):
        obs.append(f"Comparativo nacional: R$ {num(r[17]):,.2f}/ano")
    p = novo_produto(
        chave,
        nome=nome,
        fornecedorNome=forn,
        quantidadeContratada=int(num(r[4])),
        valorUnitario=num(r[6]),
        moeda=moeda,
        periodicidade="Anual" if texto(r[3]) == "Anual" else "Mensal",
        cotacao={"valor": ptax, "data": datetime.date.today().isoformat(), "fonte": "PTAX", "observacao": "PTAX implícita nos totais da planilha de licenças"} if moeda == "USD" and ptax else None,
        iofPercentual=3.5 if o else None,
        impostosPercentual=round((pimp / o - 1) * 100, 2) if o and pimp else None,
        vencimento=venc.date().isoformat() if isinstance(venc, datetime.datetime) else None,
        vencimentoTexto=texto(venc) if not isinstance(venc, datetime.datetime) else None,
        responsavel=texto(r[9]),
        utilizacao=util if not multiplo else None,
        manutencao=texto(r[10]),
        substituto=texto(r[11]),
        riscos=texto(r[12]),
        observacoes="\n".join(obs),
    )
    if p["quantidadeContratada"] == 0:
        p["confirmacoesPendentes"].append("Quantidade contratada zero na planilha")

# ------------------------------------------------------------------ Microsoft (planos próprios)
ms = list(wb["Microsoft"].iter_rows(values_only=True))
for r in ms[4:]:
    prod = texto(r[9]) if len(r) > 9 else ""
    if prod and not prod.lower().startswith("total") and num(r[10]):
        novo_produto(f"ms:{norm(prod)}", nome=prod, fornecedorNome="Microsoft", quantidadeContratada=int(num(r[10])), valorUnitario=num(r[11]), moeda="BRL", periodicidade="Mensal", responsavel="Saulo Zimbaro" if False else None)


def achar_produto(fornecedor_norm, licenca):
    ln = norm(licenca)
    cands = [p for k, p in produtos.items() if k.startswith(f"sw:{fornecedor_norm}") or k.startswith(f"{fornecedor_norm}:")]
    for p in cands:
        pn = norm(p["nome"])
        if pn[:15] and (ln.startswith(pn[:15]) or pn.startswith(ln[:15])):
            return p
    return None


com_lista = set()  # produtos com lista de usuários na planilha (uso controlado por pessoa)


def atribuir(prod, nome, email="", area="", centro="", extra=None):
    com_lista.add(prod["chave"])
    ck = registrar_colaborador(nome, email, area, centro)
    if ck:
        a = {"produtoChave": prod["chave"], "colaboradorChave": ck, "status": "Ativa"}
        if extra:
            a.update(extra)
        atribuicoes.append(a)


for r in ms[4:]:
    nome, email, lic = texto(r[1]), texto(r[2]), texto(r[4])
    if not lic or norm(email) in ("email", "e mail") or norm(lic) == "licenca" or nome.lower().startswith("total"):
        continue  # linhas de cabeçalho repetidas no meio da lista
    p = produtos.get(f"ms:{norm(lic)}") or next((p for k, p in produtos.items() if k.startswith("ms:") and norm(p["nome"]).startswith(norm(lic)[:10])), None)
    if not p:
        p = novo_produto(f"ms:{norm(lic)}", nome=lic, fornecedorNome="Microsoft", confirmacoesPendentes=["Plano sem quantidade no resumo da planilha"])
    atribuir(p, nome if nome not in ("-",) else email, email)

# ------------------------------------------------------------------ Atlassian (planos do catálogo)
for r in list(wb["Atlassian"].iter_rows(values_only=True))[4:]:
    nome, email, lic = texto(r[1]), texto(r[2]), texto(r[3])
    if not lic or nome.lower().startswith("total") or norm(email) in ("email", "e mail") or norm(lic) == "licenca":
        continue  # linhas de cabeçalho repetidas no meio da lista
    p = achar_produto("atlassian", lic) or novo_produto(f"sw:atlassian:{norm(lic)[:40]}", nome=lic, fornecedorNome="Atlassian", confirmacoesPendentes=["Plano não encontrado no catálogo da planilha"])
    atribuir(p, nome, email)

# ------------------------------------------------------------------ abas simples
ABAS = {
    "Slack": "slack", "Nuclino": "nuclino", "Scansource": "scansource", "Zoho": "zoho", "ACS Pro": "acs pro solucoes",
    "Gmail": "google cloud", "Monday": "monday", "Figma": "figma", "Cursor": "cursor", "Lucid": "lucid",
    "Figo": None, "Ahgora": None, "Wellhub": None, "Gupy": None,
}
for aba, forn in ABAS.items():
    linhas = list(wb[aba].iter_rows(values_only=True))
    cab_i = next(i for i, r in enumerate(linhas) if any(isinstance(v, str) and "Quant" in v for v in r))
    cab = [norm(c) for c in linhas[cab_i]]
    i_area = next((i for i, c in enumerate(cab) if c in ("area", "projeto")), None)
    i_email = next((i for i, c in enumerate(cab) if c in ("e mail", "email")), None)
    i_mensal = next(i for i, c in enumerate(cab) if c.startswith("valor") and "mensal" in c)
    dados = [r for r in linhas[cab_i + 1:] if any(v not in (None, "") for v in r[1:6])]
    dados = [r for r in dados if not texto(r[1]).lower().startswith("total")]
    vagas = len(dados)
    if forn:
        p = next((p for k, p in produtos.items() if k.startswith(f"sw:{forn}")), None)
    else:
        p = None
    if not p:
        unit = next((num(r[i_mensal]) for r in dados if num(r[i_mensal])), 0)
        p = novo_produto(f"aba:{norm(aba)}", nome=aba, fornecedorNome=aba, quantidadeContratada=vagas, valorUnitario=unit, moeda="BRL", periodicidade="Mensal")
        p["confirmacoesPendentes"].append("Produto criado a partir da aba (não está no catálogo Software): confirmar fornecedor, valor e vencimento")
    elif vagas and vagas != p["quantidadeContratada"]:
        p["confirmacoesPendentes"].append(f"Catálogo indica {p['quantidadeContratada']} licença(s); a aba {aba} lista {vagas} vaga(s)")
    com_lista.add(p["chave"])
    sem_nome = 0
    for r in dados:
        nome = texto(r[1])
        if not nome:
            sem_nome += 1
            continue
        if norm(nome).startswith("disponiv"):
            continue
        email = texto(r[i_email]) if i_email is not None else ""
        area = texto(r[i_area]) if i_area is not None else ""
        atribuir(p, nome, email, area=area if area.lower() not in ("akross",) else "")
    if sem_nome:
        p["confirmacoesPendentes"].append(f"{sem_nome} vaga(s) na aba {aba} sem nome de usuário")

# ------------------------------------------------------------------ telefonia (mês mais recente)
tel = [r for r in wb["Planilha1"].iter_rows(min_row=5, values_only=True) if r[2]]


def mes(r):
    m = r[28]
    if isinstance(m, datetime.datetime):
        return (m.year, m.month)
    if isinstance(m, str) and re.match(r"^\d+/\d{4}$", m):
        a, b = m.split("/")
        return (int(b), int(a))
    return None


tel = [r for r in tel if mes(r)]
ult = max(mes(r) for r in tel)
atuais = [r for r in tel if mes(r) == ult and not norm(r[3]).startswith("desativ")]
for r in atuais:
    op = texto(r[5]).strip()
    p = produtos.get(f"tel:{norm(op)}") or novo_produto(f"tel:{norm(op)}", nome=f"Linhas móveis {op}", fornecedorNome="Telefonica" if norm(op) == "vivo" else op, categoria="Linha telefônica", periodicidade="Mensal", vencimentoTexto="Mensal")
    p["quantidadeContratada"] += 1
    valor = num(r[27]) or num(r[13])
    numero = f"({texto(r[1])}) {texto(r[2])}"
    atribuir(p, texto(r[3]), "", area=texto(r[6]), centro=texto(r[7]), extra={"linha": {"numero": numero, "operadora": op, "plano": texto(r[4])}, "valorMensalBRL": round(valor, 2)})
for p in produtos.values():
    if p["categoria"] == "Linha telefônica":
        p["observacoes"] = f"Linhas da fatura de {ult[1]:02d}/{ult[0]} (Planilha1)."

# ------------------------------------------------------------------ cancelados
for r in list(wb["Cancelados"].iter_rows(min_row=2, values_only=True)):
    forn = texto(r[1])
    if not forn:
        continue
    sit = "Downgrade" if "grade" in norm(r[3]) else "Cancelado"
    novo_produto(f"canc:{norm(forn)}", nome=forn, fornecedorNome=forn, situacao=sit, utilizacao=texto(r[2]), responsavel=texto(r[10]), observacoes=texto(r[12]))

# ------------------------------------------------------------------ montagem
saida_colabs = []
for c in colabs.values():
    nomes = [n for n in c["nomes"] if n]
    melhor = max(nomes, key=lambda n: (len(n.split()), not n.isupper(), len(n))) if nomes else c["chave"][2:]
    melhor = titulo(melhor) if melhor.isupper() or melhor.islower() else melhor
    tipo = "Conta compartilhada" if COMPARTILHADA.search(norm(f"{melhor} {local(c['email'])}")) else "Pessoa"
    item = {"chave": c["chave"], "nome": melhor, "tipo": tipo, "ativo": True}
    if c["email"]:
        item["email"] = c["email"]
    if c["area"]:
        item["area"] = c["area"]
    if c["centroCusto"]:
        item["centroCusto"] = c["centroCusto"]
    if not c["email"] and tipo == "Pessoa":
        item["confirmacoesPendentes"] = ["Sem e-mail na planilha: confirmar identidade"]
    saida_colabs.append(item)

# Mesmo nome em vários e-mails: os e-mails que não são da pessoa viram contas compartilhadas com responsável.
por_nome_saida = collections.defaultdict(list)
for c in saida_colabs:
    por_nome_saida[norm(c["nome"])].append(c)
for n, grupo in por_nome_saida.items():
    if len(grupo) < 2:
        continue
    primeiro = n.split()[0] if n else ""
    pessoais = [c for c in grupo if primeiro and primeiro in local(c.get("email", ""))]
    for c in grupo:
        if c not in pessoais and c.get("email"):
            c["tipo"] = "Conta compartilhada"
            c["nome"] = f"{local(c['email'])} (responsável: {c['nome']})"
            c.pop("confirmacoesPendentes", None)
    if len(pessoais) > 1:
        for c in pessoais:
            outros = ", ".join(x["email"] for x in pessoais if x is not c and x.get("email"))
            c.setdefault("confirmacoesPendentes", []).append(f"Possível duplicidade: mesmo nome com outro e-mail ({outros})")
for c in saida_colabs:
    if c["tipo"] != "Pessoa":
        c.pop("confirmacoesPendentes", None)

# Atribuição duplicada (mesma pessoa duas vezes no mesmo produto) vira uma só.
vistos, atrib_final = set(), []
for a in atribuicoes:
    k = (a["produtoChave"], a["colaboradorChave"], (a.get("linha") or {}).get("numero"))
    if k not in vistos:
        vistos.add(k)
        atrib_final.append(a)

vinc = 0
saida_produtos = []
for p in produtos.values():
    cid, cod = contrato_sugerido(p["fornecedorNome"]) if p["situacao"] == "Ativo" else (None, None)
    if cid:
        p["contratoSugerido"] = cid
        p["confirmacoesPendentes"].append(f"Confirmar contrato vinculado automaticamente pelo nome do fornecedor ({cod})")
        vinc += 1
    if p["situacao"] == "Ativo" and p["categoria"] == "Software" and p["chave"] not in com_lista:
        p["controlePorUsuario"] = False  # sem lista de usuários na planilha: não gera custo ocioso
    if not p["confirmacoesPendentes"]:
        del p["confirmacoesPendentes"]
    saida_produtos.append({k: v for k, v in p.items() if v is not None})

json.dump({"formato": "akross-licencas-importacao", "produtos": saida_produtos, "colaboradores": saida_colabs, "atribuicoes": atrib_final}, open(out_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# ------------------------------------------------------------------ relatório
uso = collections.Counter(a["produtoChave"] for a in atrib_final)
cod_por_id = {c["contrato_id"]: c["codigo"] for c in backup["contratos"]}
w = openpyxl.Workbook()
s1 = w.active
s1.title = "Produtos"
s1.append(["Produto", "Fornecedor", "Categoria", "Situação", "Contratadas", "Atribuídas na planilha", "Contrato sugerido", "A confirmar"])
for p in saida_produtos:
    s1.append([p["nome"], p["fornecedorNome"], p["categoria"], p["situacao"], p.get("quantidadeContratada", 0), uso[p["chave"]], cod_por_id.get(p.get("contratoSugerido"), ""), "; ".join(p.get("confirmacoesPendentes", []))])
s2 = w.create_sheet("Colaboradores")
s2.append(["Nome", "E-mail", "Tipo", "Área", "Licenças", "A confirmar"])
lic_por_colab = collections.Counter(a["colaboradorChave"] for a in atrib_final)
for c in sorted(saida_colabs, key=lambda c: c["nome"]):
    s2.append([c["nome"], c.get("email", ""), c["tipo"], c.get("area", ""), lic_por_colab[c["chave"]], "; ".join(c.get("confirmacoesPendentes", []))])
for sh, larg in ((s1, [40, 26, 16, 12, 12, 14, 16, 80]), (s2, [36, 36, 20, 18, 10, 50])):
    for cell in sh[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A2C66")
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, l in enumerate(larg):
        sh.column_dimensions[chr(65 + i)].width = l
    sh.freeze_panes = "A2"
w.save(rel_path)

print(json.dumps({
    "ptax_usada": ptax,
    "produtos": len(saida_produtos),
    "ativos": sum(1 for p in saida_produtos if p["situacao"] == "Ativo"),
    "cancelados": sum(1 for p in saida_produtos if p["situacao"] != "Ativo"),
    "linhas_telefonicas": sum(1 for a in atrib_final if a.get("linha")),
    "mes_telefonia": ult,
    "colaboradores": len(saida_colabs),
    "sem_email": sum(1 for c in saida_colabs if "email" not in c),
    "contas_compartilhadas": sum(1 for c in saida_colabs if c["tipo"] != "Pessoa"),
    "atribuicoes": len(atrib_final),
    "vinculados_a_contrato": vinc,
    "sem_controle_por_pessoa": sum(1 for p in saida_produtos if p.get("controlePorUsuario") is False),
}, ensure_ascii=False, indent=1))
